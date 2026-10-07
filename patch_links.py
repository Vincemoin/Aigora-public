"""
patch_links.py

Ergaenzt jeden Chunk in der bestehenden ChromaDB-Collection um ein
"link"-Feld in den Metadaten, gematcht ueber den Dateinamen (Spalte
"Name.pdf" in RAG_Corpus.xlsx == "dateiname" in den Chunk-Metadaten).
Einmalig auszufuehren -- danach zeigen retrieve.py/gpr_core.py die Links
automatisch mit an, ohne weitere Code-Aenderung.

WICHTIG zu Seitenzahlen: die Spalte "Seiten" in RAG_Corpus.xlsx ist die
GESAMT-Seitenzahl des Dokuments (z.B. 532), keine Chunk-spezifische
Seitenangabe -- eine praezise Seitenangabe PRO ZITAT ist mit den aktuell
vorhandenen Metadaten nicht moeglich, nur ein Link auf das Gesamtdokument
(so wie von der Autor selbst als akzeptabler Fallback benannt). Dieses
Skript traegt deshalb bewusst NUR "link" ein, keine Seitenzahl.

Nutzung:
    pip install openpyxl --break-system-packages   # einmalig, falls noch nicht vorhanden
    python patch_links.py pfad/zu/RAG_Corpus.xlsx
"""

import sys
import openpyxl
import chromadb

DB_DIR = "./chroma_db"
COLLECTION_NAME = "aigora_corpus"
SHEET_NAME = "Tabellenblatt1"  # das Blatt mit der vollstaendigen Spaltenliste inkl. "Link"


def load_link_lookup(xlsx_path):
    """Liest {dateiname: link} aus der xlsx -- nur Zeilen mit BEIDEN
    Werten vorhanden werden aufgenommen, damit unvollstaendige Zeilen
    nicht versehentlich leere Links eintragen."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[SHEET_NAME]

    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    idx_dateiname = header.index("Name.pdf")
    idx_link = header.index("Link")

    lookup = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        dateiname = row[idx_dateiname]
        link = row[idx_link]
        if dateiname and link:
            lookup[dateiname] = link
    return lookup


def main():
    if len(sys.argv) < 2:
        print("Nutzung: python patch_links.py pfad/zu/RAG_Corpus.xlsx")
        return

    xlsx_path = sys.argv[1]
    lookup = load_link_lookup(xlsx_path)
    print(f"{len(lookup)} Dokument-Links aus {xlsx_path} geladen.")

    client = chromadb.PersistentClient(path=DB_DIR)
    collection = client.get_collection(COLLECTION_NAME)

    existing = collection.get(include=["metadatas"])
    ids = existing["ids"]
    metadatas = existing["metadatas"]

    updated = 0
    not_found = set()
    new_metadatas = []
    for meta in metadatas:
        dateiname = meta.get("dateiname")
        link = lookup.get(dateiname)
        new_meta = dict(meta)
        if link:
            new_meta["link"] = link
            updated += 1
        else:
            not_found.add(dateiname)
        new_metadatas.append(new_meta)

    collection.update(ids=ids, metadatas=new_metadatas)

    print(f"{updated} von {len(ids)} Chunks mit Link versehen.")
    if not_found:
        print(f"\nKein Link gefunden fuer {len(not_found)} Dokument(e) -- bitte in der xlsx pruefen:")
        for d in sorted(str(x) for x in not_found):
            print(f"  - {d}")

    print(f"\nWICHTIG: den aktualisierten '{DB_DIR}'-Ordner erneut ins Repo committen, "
          f"damit die Links auch in der deployten App ankommen.")


if __name__ == "__main__":
    main()
